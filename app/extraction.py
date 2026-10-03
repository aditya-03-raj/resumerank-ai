import os
import pdfplumber
from docx import Document

def extract_pdf(path):
    text = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text.append(page_text)

    return "\n".join(text)

def extract_docx(path):
    text = []
    document = Document(path)
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text.append(paragraph.text)

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text.append(cell.text)

    return "\n".join(text)

def extract_text(path):
    extension = os.path.splitext(path)[1].lower()

    try:
        if extension == ".pdf":
            return extract_pdf(path)
        elif extension == ".docx":
            return extract_docx(path)
        else:
            print("Unsupported file type : ",path)
            return ""
    except Exception as error:
        print("Could not read",path,"-",error)
        return ""